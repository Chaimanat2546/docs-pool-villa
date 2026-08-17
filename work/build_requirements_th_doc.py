from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(r"C:\Projects\docs-pool-villa")
OUT = ROOT / "docs" / "Poolvilla-Docs-Requirements-TH-v1.1.docx"
NAVY, BLUE, LIGHT, GRAY, TEXT = "0B2545", "2E74B5", "E8EEF5", "F2F4F7", "202124"

def font(run, size=11, bold=False, color=TEXT):
    run.font.name = "Tahoma"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Tahoma")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Tahoma")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Tahoma")
    run.font.size, run.bold, run.font.color.rgb = Pt(size), bold, RGBColor.from_string(color)

def shade(cell, fill):
    p = cell._tc.get_or_add_tcPr(); n = OxmlElement("w:shd"); n.set(qn("w:fill"), fill); p.append(n)

def margin(cell):
    p = cell._tc.get_or_add_tcPr(); m = OxmlElement("w:tcMar")
    for side, value in (("top",80),("start",120),("bottom",80),("end",120)):
        n = OxmlElement(f"w:{side}"); n.set(qn("w:w"),str(value)); n.set(qn("w:type"),"dxa"); m.append(n)
    p.append(m)

def add_h(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}"); r = p.add_run(text); font(r,{1:16,2:13,3:12}[level],True,BLUE if level<3 else NAVY)

def add_p(doc, text):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(6); p.paragraph_format.line_spacing=1.1; font(p.add_run(text))

def bullets(doc, values, numbered=False):
    for x in values:
        p=doc.add_paragraph(style="List Number" if numbered else "List Bullet"); p.paragraph_format.space_after=Pt(4); p.paragraph_format.line_spacing=1.167; font(p.add_run(x))

def table(doc, headers, rows, widths):
    t=doc.add_table(rows=1, cols=len(headers)); t.style="Table Grid"; t.alignment=WD_TABLE_ALIGNMENT.LEFT; t.autofit=False
    for i,x in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=""; c.width=Inches(widths[i]); shade(c,LIGHT); margin(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; font(c.paragraphs[0].add_run(x),10,True,NAVY)
    for row in rows:
        cells=t.add_row().cells
        for i,x in enumerate(row):
            c=cells[i]; c.text=""; c.width=Inches(widths[i]); margin(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; font(c.paragraphs[0].add_run(str(x)),9.5)
    for row in t.rows:
        for c in row.cells: c.paragraphs[0].paragraph_format.space_after=Pt(0)
    doc.add_paragraph()

def callout(doc, label, text):
    t=doc.add_table(rows=1,cols=1); t.style="Table Grid"; c=t.cell(0,0); c.text=""; shade(c,GRAY); margin(c); p=c.paragraphs[0]; font(p.add_run(label+" "),10.5,True,NAVY); font(p.add_run(text),10.5); doc.add_paragraph()

doc=Document(); s=doc.sections[0]
for a in ("top_margin","bottom_margin","left_margin","right_margin"): setattr(s,a,Inches(1))
s.header_distance=Inches(.492); s.footer_distance=Inches(.492)
normal=doc.styles["Normal"]; normal.font.name="Tahoma"; normal._element.rPr.rFonts.set(qn("w:eastAsia"),"Tahoma"); normal.font.size=Pt(11); normal.paragraph_format.space_after=Pt(6); normal.paragraph_format.line_spacing=1.1
for level,size,color,before,after in [(1,16,BLUE,16,8),(2,13,BLUE,12,6),(3,12,NAVY,8,4)]:
    st=doc.styles[f"Heading {level}"]; st.font.name="Tahoma"; st._element.rPr.rFonts.set(qn("w:eastAsia"),"Tahoma"); st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor.from_string(color); st.paragraph_format.space_before=Pt(before); st.paragraph_format.space_after=Pt(after)
h=s.header.paragraphs[0]; h.alignment=WD_ALIGN_PARAGRAPH.RIGHT; font(h.add_run("Poolvilla Documentation | เอกสารความต้องการระบบ"),8.5,False,"5B6573")
f=s.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.CENTER; font(f.add_run("รออนุมัติ | 10 สิงหาคม 2026"),8.5,False,"5B6573")
p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(12); font(p.add_run("SYSTEM REQUIREMENTS DOCUMENT"),10,True,BLUE)
p=doc.add_paragraph(); font(p.add_run("ระบบ Documentation สำหรับ Poolvilla"),26,True,NAVY)
p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(18); font(p.add_run("ขอบเขต MVP สำหรับกำหนดทิศทางผลิตภัณฑ์ ข้อมูล ความปลอดภัย และการพัฒนา"),13,False,"5B6573")
table(doc,["เวอร์ชัน","วันที่","สถานะ","ผู้รับผิดชอบ"],[["1.1","10 สิงหาคม 2026","รออนุมัติ","ทีม Poolvilla"]],[1.0,1.6,2.0,1.9])
add_h(doc,"1. วัตถุประสงค์"); add_p(doc,"กำหนดความต้องการด้านฟังก์ชัน ข้อมูล ความปลอดภัย และการทำงานของระบบ Documentation สำหรับ Poolvilla เพื่อใช้เป็นเอกสารอ้างอิงก่อนออกแบบหน้าจอ สร้างฐานข้อมูล และเริ่มพัฒนา")
add_h(doc,"2. ภาพรวมระบบ"); add_p(doc,"ระบบเป็นเว็บไซต์เอกสารแนว Mintlify มีหน้าอ่านเอกสารสาธารณะและหน้าผู้ดูแล ระบบจัดเอกสารเป็นลำดับชั้นของหมวดหมู่และบทความ ผู้ดูแลเขียนด้วย Tiptap และแทรกข้อความ รูป วิดีโอ YouTube ตาราง Code block และ Callout ได้ทุกตำแหน่ง")
callout(doc,"สถาปัตยกรรม:","Next.js + TypeScript, Supabase สำหรับ Auth/PostgreSQL/RLS และ Cloudflare Image Storage เดิมสำหรับรูปภาพ")
add_h(doc,"3. เป้าหมายและขอบเขต"); add_h(doc,"3.1 เป้าหมาย",2); bullets(doc,["ให้ผู้ใช้ค้นหา อ่าน และนำทางคู่มือ Poolvilla ได้ง่ายบนทุกขนาดหน้าจอ","อนุญาตเฉพาะผู้ใช้ที่มี role_id ตามที่กำหนดให้จัดการเอกสารได้","ใช้ Supabase project และ Cloudflare Image Storage เดิม","รองรับวงจรเนื้อหา: ร่าง, ดูตัวอย่าง, เผยแพร่, เก็บถาวร"])
add_h(doc,"3.2 สิ่งที่รวมใน MVP",2); bullets(doc,["หน้าแรก, Navigation, หน้าอ่านเอกสาร, Table of Contents, Search และ Previous/Next","หมวดหมู่แบบลำดับชั้น โดยข้อมูลรองรับหลายระดับ แต่ UI MVP จัดการ 2 ระดับ","จัดการหมวด เอกสาร ลำดับ Slug สถานะ และ Preview ใน Admin","Tiptap editor พร้อมข้อความ หัวข้อ List Link Table Code block Callout รูป และ YouTube","Media Library สำหรับดู ค้นหา ใช้ซ้ำ และจัดการรูป รวมถึงตรวจสอบ/ล้างรูปที่ไม่ได้ใช้","Supabase Auth, ตรวจสอบ Role และ RLS"])
add_h(doc,"3.3 สิ่งที่ไม่รวมใน MVP",2); bullets(doc,["ประวัติการแก้ไขและการกู้คืนเวอร์ชัน","การทำงานร่วมกัน คอมเมนต์ ขั้นตอนอนุมัติ หรือเผยแพร่ตามเวลา","Search ภายนอก Analytics dashboard หลายภาษา และ Public API"])
add_h(doc,"4. ผู้ใช้และสิทธิ์"); table(doc,["บทบาท","อ่านเอกสาร","Admin","แก้ไข"],[["ผู้เยี่ยมชม","Published เท่านั้น","ไม่ได้","ไม่ได้"],["ผู้ใช้ไม่มีสิทธิ์","Published เท่านั้น","ไม่ได้","ไม่ได้"],["ผู้ดูแล Docs","ได้","ได้","ได้"],["System service","ไม่เกี่ยวข้อง","Server เท่านั้น","ควบคุม"]],[2.15,1.55,1.6,1.2]); add_p(doc,"ให้กำหนด DOC_ADMIN_ROLE_IDS เป็นแหล่งอ้างอิงเดียวของ role_id ที่เข้าจัดการ Docs ได้ และต้องยืนยันค่าจริงจากระบบหลักก่อนพัฒนา")
sections=[
("5.1 หน้าสาธารณะ",[["FR-01","ผู้เยี่ยมชมเห็นเฉพาะเอกสารที่เผยแพร่และอยู่ในหมวดที่แสดงผล","Must"],["FR-02","หน้าแรกมีทางเข้าการค้นหาและหมวดเอกสารที่เลือกมาแสดง","Must"],["FR-03","Sidebar แสดงหมวดและเอกสารตามลำดับที่กำหนด","Must"],["FR-04","หน้าเอกสารแสดงชื่อ เนื้อหา วันที่แก้ไขล่าสุด TOC และ Previous/Next","Must"],["FR-05","สร้าง TOC จาก H2 และ H3","Must"],["FR-06","ค้นหาจากชื่อ บทสรุป และข้อความในเนื้อหาได้","Must"],["FR-07","แสดงรูปจาก Cloudflare พร้อม alt text","Must"],["FR-08","แสดง YouTube ตามตำแหน่งที่ผู้เขียนวาง","Must"]]),
("5.2 สิทธิ์ Admin",[["FR-09","ต้องเข้าสู่ระบบก่อนเข้าถึง /admin","Must"],["FR-10","ตรวจสอบบทบาทก่อนแสดงฟังก์ชันผู้ดูแล","Must"],["FR-11","Supabase RLS บังคับสิทธิ์แยกจาก UI","Must"],["FR-12","ผู้ไม่มีสิทธิ์อ่าน Draft/Archived ผ่าน API ไม่ได้","Must"]]),
("5.3 จัดการหมวดหมู่",[["FR-13","สร้าง เปลี่ยนชื่อ ซ่อน/แสดง จัดลำดับ และลบหมวดได้","Must"],["FR-14","สร้างหมวดหมู่ย่อยได้","Must"],["FR-15","ข้อมูลใช้ parent_id แบบ Recursive; UI MVP จัดการ 2 ระดับ","Must"],["FR-16","ลบหมวดไม่ได้หากยังมีหมวดลูกหรือเอกสาร","Must"]]),
("5.4 จัดการเอกสาร",[["FR-17","สร้าง แก้ไข Preview Publish กลับเป็น Draft Archive จัดลำดับ ย้าย และลบได้","Must"],["FR-18","เอกสารมีสถานะ draft, published, archived","Must"],["FR-19","Publish บันทึก published_at และเนื้อหาที่ไม่ Published ห้ามเป็นสาธารณะ","Must"],["FR-20","เอกสารมี title, section, slug, excerpt, Tiptap content และ sort order","Must"],["FR-21","ตรวจ Slug และป้องกัน URL ซ้ำ","Must"],["FR-22","Preview เฉพาะผู้ดูแลและใช้ Viewer style เดียวกับหน้าสาธารณะ","Must"]]),
("5.5 Editor และสื่อ",[["FR-23","เก็บเนื้อหาเป็น Tiptap JSON ที่ผ่าน Schema validation","Must"],["FR-24","รองรับ Heading, Paragraph, List, Link, Table, Code, Callout, Image และ YouTube","Must"],["FR-25","มี Toolbar และ Slash Command","Should"],["FR-26","Editor และ Viewer ใช้ Node styling ชุดเดียวกัน","Must"],["FR-27","Upload ผ่าน Server endpoint ที่ตรวจสิทธิ์และส่งกลับ Cloudflare URL/Image ID","Must"],["FR-28","ตรวจ URL YouTube และแทรกหลายวิดีโอได้ทุกตำแหน่ง","Must"],["FR-29","ดู ค้นหา และใช้รูปจาก Media Library ซ้ำในหลายเอกสารได้","Must"],["FR-30","แสดงตัวอย่างรูป ชื่อไฟล์ URL/Image ID ขนาด วันที่อัปโหลด และผู้อัปโหลด","Must"],["FR-31","ตรวจจำนวนเอกสารที่อ้างอิงรูปแต่ละรายการได้","Must"],["FR-32","ลบรูปได้เมื่อไม่มีการอ้างอิง หรือยืนยันผลกระทบก่อนลบ","Must"],["FR-33","ดูและล้างรูปที่ไม่มีการอ้างอิงเป็นรายรายการหรือหลายรายการได้","Must"],["FR-34","ล้างรูปในระบบและ Cloudflare ให้สอดคล้อง หรือบันทึก failed เพื่อลองใหม่","Must"]])]
add_h(doc,"5. Functional Requirements")
for title, rows in sections: add_h(doc,title,2); table(doc,["ID","ความต้องการ","ระดับ"],rows,[.7,4.85,.95])
add_h(doc,"6. Information Architecture และ Route"); table(doc,["ส่วน","Route","หน้าที่"],[["หน้าแรก","/","จุดเริ่มต้นค้นหาและเลือกเอกสาร"],["ค้นหา","/search","ค้นหาเอกสาร Published"],["หน้าเอกสาร","/[...slug]","แสดงเอกสารตามเส้นทาง"],["ภาพรวม Admin","/admin","สถานะเอกสารและรายการล่าสุด"],["จัดการเอกสาร","/admin/documents","ค้นหา กรอง และจัดการ"],["สร้างเอกสาร","/admin/documents/new","สร้างเอกสารใหม่"],["แก้ไขเอกสาร","/admin/documents/[id]/edit","แก้ไข Preview และ Publish"],["โครงสร้าง","/admin/structure","จัดการ Tree และลำดับหมวด"]],[1.25,2.15,3.1])
add_h(doc,"7. ความต้องการข้อมูล"); add_h(doc,"7.1 doc_sections",2); table(doc,["Field","Type","กติกา"],[["id","UUID","Primary key"],["parent_id","UUID, nullable","อ้างถึง doc_sections.id; null คือหมวดบนสุด"],["name","text","ต้องมี"],["slug","text","ต้องมีและใช้ใน URL ได้"],["sort_order","integer","ลำดับระดับเดียวกัน"],["is_published","boolean","ควบคุมการแสดงผลสาธารณะ"],["created_by / updated_by","UUID","ID ผู้ใช้ระบบเดิม"],["created_at / updated_at","timestamptz","เวลาที่จัดการโดย Server"]],[1.85,1.5,3.15]); add_h(doc,"7.2 documents",2); table(doc,["Field","Type","กติกา"],[["id","UUID","Primary key"],["section_id","UUID","อ้างถึง doc_sections.id"],["title","text","ต้องมี"],["slug","text","ต้องมีและไม่ซ้ำในเส้นทาง"],["excerpt","text, nullable","คำเกริ่นผลค้นหา"],["content","JSONB","Tiptap JSON ที่ถูกต้อง"],["search_text","text","Plain text สำหรับค้นหา"],["status","text/enum","draft, published, archived"],["sort_order","integer","ลำดับในหมวด"],["created_by / updated_by","UUID","ID ผู้ใช้ระบบเดิม"],["published_at","timestamptz, nullable","บันทึกเมื่อเผยแพร่"],["created_at / updated_at","timestamptz","เวลาที่จัดการโดย Server"]],[1.85,1.5,3.15])
add_h(doc,"8. Business Rules"); bullets(doc,["เอกสารเป็นสาธารณะได้เมื่อ Published และทุกหมวดในเส้นทางแสดงผลได้","ห้ามเปิดเผย Draft หรือ Archived ผ่าน URL ตรง","ใช้ sort_order เพื่อกำหนดลำดับที่แน่นอน","Slug ต้องเป็นตัวพิมพ์เล็ก ใช้ใน URL ได้ และไม่ว่าง; หัวข้อไทยกำหนด slug เองได้","ตรวจ Tiptap schema ก่อนบันทึกและแสดงผล; ห้าม Render HTML ที่ไม่น่าเชื่อถือโดยตรง","JSON เก็บเฉพาะ Cloudflare URL/Image ID และ alt text","Cloudflare credential สำหรับ upload อยู่ฝั่ง Server เท่านั้น","ก่อนลบรูป ระบบต้องตรวจการอ้างอิงจาก Tiptap JSON ทั้งหมด","งานล้างรูปที่ไม่ได้ใช้เป็นการกระทำของผู้ดูแล และต้องรายงานรายการที่สำเร็จ/ไม่สำเร็จ"])
add_h(doc,"9. ความปลอดภัย"); bullets(doc,["เปิด RLS สำหรับตารางใหม่ทุกตารางที่เข้าถึงผ่าน Data API","Public SELECT ได้เฉพาะ Published และเส้นทางหมวดที่มองเห็น","การเขียนและการอ่านข้อมูลไม่สาธารณะต้องมี role_id ใน DOC_ADMIN_ROLE_IDS","Role lookup ห้ามเปิดโอกาสให้ยกระดับสิทธิ์ตนเอง","ห้ามส่ง service-role key หรือ Cloudflare credential ไป Client","ป้องกัน Admin route ฝั่ง Server และย้ำด้วย RLS","ตรวจชนิด ขนาด และมิติของรูป รวมถึง Origin YouTube","เก็บ actor id และเวลาเมื่อผู้ดูแลเปลี่ยนข้อมูล"])
add_h(doc,"10. Search"); bullets(doc,["MVP ใช้ PostgreSQL Full-text Search จาก title, excerpt และ search_text","สกัด search_text จาก Tiptap JSON เมื่อบันทึกหรือ Publish","ผลค้นหาแสดงเฉพาะเอกสาร Published ที่เข้าถึงสาธารณะได้","ไม่ค้นหา JSONB โดยตรงเป็นวิธีหลัก"])
add_h(doc,"11. Non-functional Requirements"); table(doc,["ID","ความต้องการ"],[["NFR-01","รองรับ mobile, tablet และ desktop"],["NFR-02","URL แชร์ได้ คงที่ และเหมาะกับ SEO"],["NFR-03","รองรับภาษาไทยใน heading, search และ slug"],["NFR-04","Reader ใช้ server rendering หรือ static regeneration ตาม deployment"],["NFR-05","ใช้ Cloudflare optimisation/delivery เท่าที่ storage เดิมรองรับ"],["NFR-06","Error และ upload failure แก้ได้โดยไม่เปิดเผย credential"],["NFR-07","Accessibility ขั้นพื้นฐาน: semantic heading, keyboard, focus, label และ alt text"]],[.85,5.65])
add_h(doc,"12. เกณฑ์ยอมรับ MVP"); bullets(doc,["ผู้เยี่ยมชมเปิดและค้นหาได้เฉพาะคู่มือ Published","ผู้ไม่มีสิทธิ์เข้า Admin หรือแก้ข้อมูลผ่าน Supabase ไม่ได้","ผู้ดูแลสร้าง Draft ใส่ข้อความ รูป Callout และ YouTube หลายรายการ แล้ว Preview/Publish ได้","หน้าที่ Publish แสดงเนื้อหาและรูปแบบเดียวกับ Preview","สร้างหมวดหลัก หมวดย่อย เอกสาร และจัดลำดับได้ โดย Sidebar ต้องตรงตามลำดับ","Draft/Archived ไม่ปรากฏใน Navigation, Search หรือ URL สาธารณะ","รูปอัปโหลดผ่าน Cloudflare service เดิมและ URL ถูกเก็บใน JSON","Search หาเอกสาร Published เจอจาก title, excerpt หรือเนื้อหา","ผู้ดูแลค้นหารูปจาก Media Library ใช้รูปเดิมในเอกสาร และเห็นจำนวนการอ้างอิงได้","ระบบไม่ลบรูปที่ยังถูกอ้างอิงโดยไม่ยืนยัน และล้างรูปที่ไม่มีการอ้างอิงได้"])
add_h(doc,"13. สิ่งที่ต้องยืนยันก่อนพัฒนา"); table(doc,["หัวข้อ","คำตอบที่ต้องการ","ผลกระทบ"],[["สิทธิ์ผู้ดูแล","role_id ใดจัดการ Docs ได้","RLS และ Route guard"],["ตารางผู้ใช้","ตาราง/view ใดมี role_id และเชื่อม auth.users อย่างไร","RLS"],["Cloudflare Storage","Cloudflare Images, R2 หรือ Custom API","Upload endpoint และ URL"],["Public domain","ยืนยัน domain/subdomain","Routing, Cookie, SEO"],["กติกา Slug","ไม่ซ้ำทั้งระบบหรือเฉพาะเส้นทาง","Constraint และ lookup"],["การลบ","Hard delete, Soft delete หรือ Archive only","Retention และ UI"],["Thai Search","ความคาดหวังเรื่องตัดคำไทย","PostgreSQL configuration"]],[1.3,2.85,2.35])
add_h(doc,"14. แนวทางการพัฒนาแบบแยก Module"); add_p(doc,"ไม่พัฒนาทุก Module พร้อมกัน แต่ทำให้จบเป็นส่วน ๆ โดยแต่ละ Module ต้องมี Migration/Policy, UI, Test และเกณฑ์ยอมรับของตัวเองก่อนเริ่มส่วนถัดไป")
table(doc,["ลำดับ","Module","ผลลัพธ์ที่ต้องจบในรอบนั้น"],[["M01","Foundation & Authorization","เชื่อม Supabase, role mapping, Admin route guard, RLS และ schema หลัก"],["M02","Structure Management","CRUD หมวดหลัก/ย่อย, ลำดับ และกติกาการลบ"],["M03","Document Management","CRUD เอกสาร, Draft/Publish/Archive, Slug, Preview และลำดับ"],["M04","Public Documentation","หน้าอ่าน, Sidebar, TOC, Previous/Next และการซ่อนเนื้อหาที่ไม่ Published"],["M05","Editor & Media Upload","Tiptap, Schema validation, Upload Cloudflare และ YouTube embed"],["M06","Media Library & Cleanup","รายการ/ค้นหารูป, ใช้ซ้ำ, ตรวจการอ้างอิง และล้างรูปที่ไม่ใช้"],["M07","Search & Hardening","Full-text search, Accessibility, Security test และทดสอบเกณฑ์ยอมรับรวม"]],[.65,1.65,4.2]); add_p(doc,"ก่อนขึ้น Module ใหม่ ให้ทดสอบ acceptance criteria ของ Module ก่อนหน้าและแก้ defect ที่เป็น blocker ให้จบก่อน")
OUT.parent.mkdir(parents=True,exist_ok=True); doc.save(OUT); print(OUT)
