from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(r"C:\Projects\docs-pool-villa")
OUT = ROOT / "docs" / "Poolvilla-Docs-Requirements-v1.docx"

NAVY = "0B2545"
BLUE = "2E74B5"
LIGHT = "E8EEF5"
GRAY = "F2F4F7"
TEXT = "202124"

def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)

def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")

def set_font(run, size=11, bold=False, color=TEXT):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)

def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.style = f"Heading {level}"
    r = p.add_run(text)
    set_font(r, {1:16, 2:13, 3:12}[level], True, BLUE if level < 3 else NAVY)
    return p

def add_para(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.10
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix); set_font(r, 11, True)
        r = p.add_run(text[len(bold_prefix):]); set_font(r)
    else:
        r = p.add_run(text); set_font(r)
    return p

def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.167
        r = p.add_run(item); set_font(r)

def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    table.autofit = False
    for i, head in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        set_cell_shading(cell, LIGHT)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        r = cell.paragraphs[0].add_run(head); set_font(r, 10, True, NAVY)
        if widths: cell.width = Inches(widths[i])
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = ""
            set_cell_margins(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            r = cells[i].paragraphs[0].add_run(str(value)); set_font(r, 9.5)
            if widths: cells[i].width = Inches(widths[i])
    for row in table.rows:
        for cell in row.cells:
            cell.paragraphs[0].paragraph_format.space_after = Pt(0)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table

def add_callout(doc, label, text):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    cell.text = ""
    set_cell_shading(cell, GRAY); set_cell_margins(cell, 120, 160, 120, 160)
    p = cell.paragraphs[0]
    r = p.add_run(label + " "); set_font(r, 10.5, True, NAVY)
    r = p.add_run(text); set_font(r, 10.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(1); section.bottom_margin = Inches(1)
section.left_margin = Inches(1); section.right_margin = Inches(1)
section.header_distance = Inches(0.492); section.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"; normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri"); normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
normal.font.size = Pt(11); normal.font.color.rgb = RGBColor.from_string(TEXT)
normal.paragraph_format.space_after = Pt(6); normal.paragraph_format.line_spacing = 1.10
for level, size, color, before, after in [(1,16,BLUE,16,8),(2,13,BLUE,12,6),(3,12,NAVY,8,4)]:
    s = styles[f"Heading {level}"]
    s.font.name = "Calibri"; s._element.rPr.rFonts.set(qn("w:ascii"), "Calibri"); s._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    s.font.size = Pt(size); s.font.bold = True; s.font.color.rgb = RGBColor.from_string(color)
    s.paragraph_format.space_before = Pt(before); s.paragraph_format.space_after = Pt(after)

header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = header.add_run("Poolvilla Documentation Platform | System Requirements")
set_font(r, 8.5, False, "5B6573")
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = footer.add_run("Draft for approval | 10 August 2026")
set_font(r, 8.5, False, "5B6573")

p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(12); p.paragraph_format.space_after = Pt(4)
r = p.add_run("SYSTEM REQUIREMENTS DOCUMENT"); set_font(r, 10, True, BLUE)
p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(6)
r = p.add_run("Poolvilla Documentation Platform"); set_font(r, 26, True, NAVY)
p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(18)
r = p.add_run("MVP baseline for product direction, security, data design, and delivery scope"); set_font(r, 13, False, "5B6573")
add_table(doc, ["Version", "Date", "Status", "Owner"], [["1.0", "10 August 2026", "Draft for approval", "Poolvilla Team"]], [1.0, 1.6, 2.0, 1.9])

add_heading(doc, "1. Purpose")
add_para(doc, "This document defines the functional, data, security, and operational requirements for the Poolvilla documentation platform. It is the baseline to approve before UI design, database migration, and implementation begin.")

add_heading(doc, "2. Product Summary")
add_para(doc, "The product is a Mintlify-inspired documentation website with a responsive public reader experience and a restricted admin workspace. Documentation is organised as a hierarchy of sections and documents. Authors write with Tiptap and can place formatted text, images, YouTube videos, tables, code blocks, and callouts anywhere in an article.")
add_callout(doc, "Architecture:", "Next.js + TypeScript, Supabase for Auth/PostgreSQL/RLS, and the existing Cloudflare image storage for media delivery.")

add_heading(doc, "3. Goals and Scope")
add_heading(doc, "3.1 Goals", 2)
add_bullets(doc, ["Make Poolvilla guides easy to find, read, and navigate on desktop and mobile.", "Allow only users with approved existing role_id values to manage documentation.", "Reuse the existing Supabase project and Cloudflare image storage.", "Support a safe lifecycle: draft, preview, publish, and archive."])
add_heading(doc, "3.2 Included in MVP", 2)
add_bullets(doc, ["Public home, navigation, document view, table of contents, search, and previous/next navigation.", "Hierarchical sections (the data model supports any depth; the MVP admin UI starts with two levels).", "Admin management of sections, documents, ordering, slugs, publication state, and previews.", "Tiptap editor with text, headings, lists, links, tables, code blocks, callouts, images, and YouTube embeds.", "Supabase authentication, role checks, and Row Level Security (RLS)."])
add_heading(doc, "3.3 Excluded from MVP", 2)
add_bullets(doc, ["Document version history and restore.", "Collaboration, commenting, review approvals, or scheduled publishing.", "Full media-library management and unused-media cleanup.", "External search engines, analytics dashboard, multilingual content, and public API."])

add_heading(doc, "4. Users and Permissions")
add_table(doc, ["Role", "Public docs", "Admin", "Write"], [["Guest", "Published only", "No", "No"], ["Authenticated user without approved role", "Published only", "No", "No"], ["Docs administrator", "Yes", "Yes", "Yes"], ["System service", "N/A", "Server-side integration only", "Controlled"]], [2.15, 1.55, 1.6, 1.2])
add_para(doc, "DOC_ADMIN_ROLE_IDS is the configuration source for allowed role_id values. Exact values must be confirmed from the main Poolvilla system before implementation.")

add_heading(doc, "5. Functional Requirements")
for title, headers, rows, widths in [
    ("5.1 Public documentation", ["ID", "Requirement", "Priority"], [["FR-01", "Visitors can see only published and visible documentation.", "Must"], ["FR-02", "Home page provides a search entry point and documentation discovery.", "Must"], ["FR-03", "Sidebar shows sections and documents in configured order.", "Must"], ["FR-04", "Document page shows title, content, last updated, TOC, and previous/next links.", "Must"], ["FR-05", "TOC is generated from H2 and H3 headings.", "Must"], ["FR-06", "Search returns published documents matching title, excerpt, or body text.", "Must"], ["FR-07", "Cloudflare-hosted images render with supplied alt text.", "Must"], ["FR-08", "YouTube embeds render in their authored position.", "Must"]], [0.7, 4.85, 0.95]),
    ("5.2 Admin access", ["ID", "Requirement", "Priority"], [["FR-09", "Users sign in before accessing /admin.", "Must"], ["FR-10", "Application checks approved role before rendering admin functions.", "Must"], ["FR-11", "Supabase RLS separately enforces database access.", "Must"], ["FR-12", "Unauthorised users cannot read drafts or archived content via APIs.", "Must"]], [0.7, 4.85, 0.95]),
    ("5.3 Section management", ["ID", "Requirement", "Priority"], [["FR-13", "Admins can create, rename, describe, hide/show, reorder, and delete sections.", "Must"], ["FR-14", "Admins can create child sections.", "Must"], ["FR-15", "Data uses recursive parent_id; MVP UI manages two visible levels.", "Must"], ["FR-16", "Deletion is blocked until child sections/documents are moved or removed.", "Must"]], [0.7, 4.85, 0.95]),
    ("5.4 Document management", ["ID", "Requirement", "Priority"], [["FR-17", "Admins can create, edit, preview, publish, draft, archive, reorder, move, and delete documents.", "Must"], ["FR-18", "Document state is draft, published, or archived.", "Must"], ["FR-19", "Publishing records published_at; non-published documents are not public.", "Must"], ["FR-20", "Document contains title, section, slug, optional excerpt, Tiptap content, and sort order.", "Must"], ["FR-21", "Slug validation prevents duplicate public routes.", "Must"], ["FR-22", "Admin-only preview uses the same viewer style as public pages.", "Must"]], [0.7, 4.85, 0.95]),
    ("5.5 Editor and media", ["ID", "Requirement", "Priority"], [["FR-23", "Editor stores document body as valid Tiptap JSON.", "Must"], ["FR-24", "Supported blocks: headings, text, lists, link, table, code, quote, divider, callout, image, YouTube.", "Must"], ["FR-25", "Toolbar and slash-command menu expose supported blocks.", "Should"], ["FR-26", "Editor and viewer share node styling, except for editing controls.", "Must"], ["FR-27", "Server upload endpoint authorises and returns a Cloudflare URL or image ID.", "Must"], ["FR-28", "YouTube URLs are validated and may be inserted multiple times anywhere in content.", "Must"]], [0.7, 4.85, 0.95]),
]:
    add_heading(doc, title, 2); add_table(doc, headers, rows, widths)

add_heading(doc, "6. Information Architecture and Routes")
add_table(doc, ["Area", "Route", "Purpose"], [["Public home", "/", "Search entry point and documentation discovery"], ["Public search", "/search", "Search published documentation"], ["Public document", "/[...slug]", "Render document by hierarchical route"], ["Admin overview", "/admin", "Document status and recently updated items"], ["Documents", "/admin/documents", "Filter and manage documents"], ["New document", "/admin/documents/new", "Create a document"], ["Edit document", "/admin/documents/[id]/edit", "Edit, preview, and publish"], ["Structure", "/admin/structure", "Manage section tree and order"]], [1.25, 2.15, 3.1])

add_heading(doc, "7. Data Requirements")
add_heading(doc, "7.1 doc_sections", 2)
add_table(doc, ["Field", "Type", "Rule"], [["id", "UUID", "Primary key"], ["parent_id", "UUID, nullable", "References doc_sections.id; null for top level"], ["name", "text", "Required"], ["slug", "text", "Required and route-safe"], ["description", "text, nullable", "Optional summary"], ["sort_order", "integer", "Required sibling order"], ["is_published", "boolean", "Controls public visibility"], ["created_by / updated_by", "UUID", "Existing authenticated user id"], ["created_at / updated_at", "timestamptz", "Server-managed timestamps"]], [1.85, 1.5, 3.15])
add_heading(doc, "7.2 documents", 2)
add_table(doc, ["Field", "Type", "Rule"], [["id", "UUID", "Primary key"], ["section_id", "UUID", "References doc_sections.id"], ["title", "text", "Required"], ["slug", "text", "Required; unique in resolved public route"], ["excerpt", "text, nullable", "Optional result summary"], ["content", "JSONB", "Required valid Tiptap JSON"], ["search_text", "text", "Derived plain text for search"], ["status", "text/enum", "draft, published, archived"], ["sort_order", "integer", "Required document order"], ["created_by / updated_by", "UUID", "Existing authenticated user id"], ["published_at", "timestamptz, nullable", "Set when published"], ["created_at / updated_at", "timestamptz", "Server-managed timestamps"]], [1.85, 1.5, 3.15])

add_heading(doc, "8. Business Rules")
add_bullets(doc, ["Only a published document under a publicly visible ancestor path is publicly accessible.", "Direct access to draft or archived content must not expose its content.", "Sibling order is deterministic through sort_order.", "A slug is lowercase, URL-safe, and non-empty; Thai titles may use a manual or transliterated slug.", "Content must validate against the supported Tiptap schema before storage and rendering.", "Untrusted HTML must not be rendered directly; node attributes and links are sanitised.", "Document JSON stores only Cloudflare delivery URL/identifier and alt text; media is not stored in Supabase.", "Cloudflare upload credentials are server-side only."])

add_heading(doc, "9. Security Requirements")
add_bullets(doc, ["Enable RLS on all new public tables exposed through Supabase Data API.", "Public SELECT policies permit only published documents and visible section paths.", "All writes and non-public reads require an authenticated user with role_id in DOC_ADMIN_ROLE_IDS.", "Role lookup must not permit users to assign or elevate their own role.", "Never expose Supabase service-role or Cloudflare upload credentials in client code.", "Protect admin routes server-side, then mirror the same rule in RLS.", "Validate upload file type, size, dimensions where needed, filename handling, and allowed YouTube origin.", "Log meaningful admin mutations with actor id and timestamps at minimum."])

add_heading(doc, "10. Search Requirements")
add_bullets(doc, ["MVP search uses PostgreSQL full-text search over title, excerpt, and search_text.", "search_text is extracted from Tiptap JSON when the document is saved or published.", "Search returns only publicly visible published documents.", "Result items show title, excerpt or summary, and hierarchy path.", "JSONB is not queried directly as the primary search mechanism."])

add_heading(doc, "11. Non-functional Requirements")
add_table(doc, ["ID", "Requirement"], [["NFR-01", "Public pages are responsive for mobile, tablet, and desktop."], ["NFR-02", "Public URLs are shareable, stable, and SEO-friendly."], ["NFR-03", "Thai text works correctly in headings, search, and administrator-selected slugs."], ["NFR-04", "Reader pages are server-rendered or statically regenerated where compatible with deployment."], ["NFR-05", "Existing Cloudflare optimisation/delivery capabilities are used where available."], ["NFR-06", "Admin errors and failed uploads are actionable without exposing internals or credentials."], ["NFR-07", "Baseline accessibility: semantic headings, keyboard navigation, focus states, labelled controls, and meaningful alt text."]], [0.85, 5.65])

add_heading(doc, "12. MVP Acceptance Criteria")
add_bullets(doc, ["A guest can navigate and search only published guides.", "A user without an approved role cannot access admin UI or mutate records through Supabase.", "An administrator can create a draft, add text, image, callout, and multiple YouTube embeds, preview, and publish it.", "Published output matches the preview content order and viewer styling.", "Administrators can add a parent section, child section, and documents, then reorder them; the sidebar follows that order.", "Draft and archived documents do not appear in public navigation, search, or direct public requests.", "An uploaded image reaches the existing Cloudflare service and its returned URL is saved in document JSON.", "Search finds a published document by words in title, excerpt, or body."])

add_heading(doc, "13. Open Decisions Before Implementation")
add_table(doc, ["Decision", "Needed answer", "Impact"], [["Admin roles", "Which existing role_id values administer docs?", "RLS and route guards"], ["User table", "Which table/view owns role_id and links to auth.users?", "RLS implementation"], ["Cloudflare storage", "Cloudflare Images, R2, or custom upload API?", "Upload endpoint and delivery URLs"], ["Public host", "Confirm production domain/subdomain.", "Routing, cookies, SEO"], ["Slug rule", "Globally unique or unique only within a section path?", "Constraints and route lookup"], ["Deletion policy", "Hard delete, soft delete, or archive-only?", "Retention and UI"], ["Thai search", "Expected word-segmentation quality?", "PostgreSQL search configuration"]], [1.3, 2.85, 2.35])

add_heading(doc, "14. Recommended Delivery Order")
for i, item in enumerate(["Confirm open decisions, data ownership, and role mapping.", "Create migrations, constraints, indexes, and RLS policies.", "Build authentication helpers, admin route guard, and section/document server actions.", "Build public viewer, hierarchy navigation, and search.", "Build Tiptap editor, validation, preview, and Cloudflare upload integration.", "Build admin document/structure screens and ordering.", "Test authorisation, public visibility, media upload, and acceptance criteria."], 1):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(item); set_font(r)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
