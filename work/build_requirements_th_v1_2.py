from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Projects\docs-pool-villa")
SOURCE = ROOT / "docs" / "Poolvilla-Docs-Requirements-TH-v1.2.md"
OUTPUT = ROOT / "docs" / "Poolvilla-Docs-Requirements-TH-v1.2.docx"

# compact_reference_guide preset + named override `thai_font = Tahoma`
FONT = "Tahoma"
CODE_FONT = "Consolas"
BLUE = "2E74B5"
NAVY = "1F4D78"
INK = "202124"
MUTED = "5B6573"
TABLE_HEADER = "E8EEF5"
TABLE_ALT = "F7F9FB"
BORDER = "C7D0DA"
CONTENT_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGINS = {"top": 80, "bottom": 80, "start": 120, "end": 120}


def set_font(run, *, name=FONT, size=11, bold=None, italic=None, color=INK):
    run.font.name = name
    run._element.get_or_add_rPr()
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    old = tc_pr.find(qn("w:tcMar"))
    if old is not None:
        tc_pr.remove(old)
    tc_mar = OxmlElement("w:tcMar")
    for side, value in CELL_MARGINS.items():
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        tc_mar.append(node)
    tc_pr.append(tc_mar)


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    if sum(widths) != CONTENT_DXA:
        raise ValueError(f"Table widths must total {CONTENT_DXA}: {widths}")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[index])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def choose_widths(headers):
    count = len(headers)
    if count == 2:
        if headers[0].lower() in {"field", "ส่วน", "รายการ", "บทบาท"}:
            return [2600, 6760]
        return [3300, 6060]
    if count == 3:
        if headers[0].lower() in {"field", "id"}:
            return [1900, 1900, 5560]
        return [1800, 3780, 3780]
    if count == 4:
        return [2000, 2450, 2450, 2460]
    base = CONTENT_DXA // count
    widths = [base] * count
    widths[-1] += CONTENT_DXA - sum(widths)
    return widths


def add_inline(paragraph, text, *, default_size=11, default_color=INK, default_bold=False):
    pattern = re.compile(r"(\*\*.+?\*\*|`.+?`)")
    cursor = 0
    for match in pattern.finditer(text):
        if match.start() > cursor:
            set_font(
                paragraph.add_run(text[cursor : match.start()]),
                size=default_size,
                color=default_color,
                bold=default_bold,
            )
        token = match.group(0)
        if token.startswith("**"):
            set_font(
                paragraph.add_run(token[2:-2]),
                size=default_size,
                color=default_color,
                bold=True,
            )
        else:
            run = paragraph.add_run(token[1:-1])
            set_font(run, name=CODE_FONT, size=max(default_size - 1, 8.5), color=NAVY)
            run.font.highlight_color = None
        cursor = match.end()
    if cursor < len(text):
        set_font(
            paragraph.add_run(text[cursor:]),
            size=default_size,
            color=default_color,
            bold=default_bold,
        )


def add_table(doc, headers, rows):
    widths = choose_widths(headers)
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"

    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = ""
        set_cell_shading(cell, TABLE_HEADER)
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        add_inline(paragraph, header, default_size=9.5, default_color=NAVY, default_bold=True)
    set_repeat_table_header(table.rows[0])

    for row_index, values in enumerate(rows, start=1):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cell = cells[index]
            cell.text = ""
            if row_index % 2 == 0:
                set_cell_shading(cell, TABLE_ALT)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.15
            add_inline(paragraph, value, default_size=9.25)

    set_table_geometry(table, widths)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def clean_table_cell(text):
    return text.strip().strip("|").strip()


def parse_table(lines, start):
    headers = [clean_table_cell(item) for item in lines[start].split("|")[1:-1]]
    rows = []
    index = start + 2
    while index < len(lines) and lines[index].strip().startswith("|"):
        rows.append([clean_table_cell(item) for item in lines[index].split("|")[1:-1]])
        index += 1
    return headers, rows, index


def add_list_item(doc, text, numbered=False):
    style = "List Number" if numbered else "List Bullet"
    paragraph = doc.add_paragraph(style=style)
    paragraph.paragraph_format.left_indent = Inches(0.375)
    paragraph.paragraph_format.first_line_indent = Inches(-0.188)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.25
    add_inline(paragraph, text)


def add_heading(doc, text, level):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    add_inline(
        paragraph,
        text,
        default_size={1: 16, 2: 13, 3: 12}[level],
        default_color=BLUE if level < 3 else NAVY,
        default_bold=True,
    )


def add_body(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.25
    add_inline(paragraph, text)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(paragraph.add_run("Poolvilla Docs Requirements  |  "), size=8.5, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    paragraph._p.append(begin)
    paragraph._p.append(instr)
    paragraph._p.append(end)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading_tokens = {
        1: (16, BLUE, 18, 10),
        2: (13, BLUE, 14, 7),
        3: (12, NAVY, 10, 5),
    }
    for level, (size, color, before, after) in heading_tokens.items():
        style = doc.styles[f"Heading {level}"]
        style.font.name = FONT
        style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(
        header.add_run("Poolvilla Documentation  |  System Requirements v1.2"),
        size=8.5,
        color=MUTED,
    )
    add_page_number(section.footer.paragraphs[0])


def add_masthead(doc):
    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_before = Pt(12)
    kicker.paragraph_format.space_after = Pt(4)
    set_font(kicker.add_run("SYSTEM REQUIREMENTS DOCUMENT"), size=10, bold=True, color=BLUE)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(6)
    set_font(title.add_run("ระบบ Documentation สำหรับ Poolvilla"), size=26, bold=True, color=NAVY)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(16)
    set_font(
        subtitle.add_run("Requirement baseline สำหรับ MVP และการพัฒนาแบบแยก Module"),
        size=13,
        color=MUTED,
    )

    add_table(
        doc,
        ["เวอร์ชัน", "วันที่", "สถานะ", "ผู้รับผิดชอบ"],
        [["1.2", "10 สิงหาคม 2026", "ยืนยันแล้ว — รอเริ่ม M01", "ทีม Poolvilla"]],
    )


def build():
    text = SOURCE.read_text(encoding="utf-8")
    lines = text.splitlines()
    start = next(i for i, line in enumerate(lines) if line.startswith("## 1."))

    doc = Document()
    configure_document(doc)
    add_masthead(doc)

    index = start
    while index < len(lines):
        raw = lines[index].rstrip()
        stripped = raw.strip()
        if not stripped:
            index += 1
            continue
        if stripped.startswith("### "):
            add_heading(doc, stripped[4:], 2)
            index += 1
            continue
        if stripped.startswith("## "):
            add_heading(doc, stripped[3:], 1)
            index += 1
            continue
        if stripped.startswith("#### "):
            add_heading(doc, stripped[5:], 3)
            index += 1
            continue
        if stripped.startswith("|") and index + 1 < len(lines) and re.match(r"^\|[-:| ]+\|$", lines[index + 1].strip()):
            headers, rows, index = parse_table(lines, index)
            add_table(doc, headers, rows)
            continue
        if stripped.startswith("- "):
            add_list_item(doc, stripped[2:])
            index += 1
            continue
        number_match = re.match(r"^\d+\.\s+(.*)$", stripped)
        if number_match:
            add_list_item(doc, number_match.group(1), numbered=True)
            index += 1
            continue
        add_body(doc, stripped.replace("  ", " "))
        index += 1

    doc.core_properties.title = "Poolvilla Docs System Requirements v1.2"
    doc.core_properties.subject = "MVP requirements and module delivery rules"
    doc.core_properties.author = "Poolvilla"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
